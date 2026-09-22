from pathlib import Path
from openpyxl import load_workbook
from app.parsers.base import Parser
from app.ingestion.result import IngestionResult,RecordError
EXPECTED=['CENTRO DE COSTO', 'CUENTA', 'DIVISA', 'CONTRATO', 'REFERENCIA DE CRUCE', 'IMPORTE', 'DESCRIPCION', 'FECHA', 'TIPO DE DOCUMENTO', 'NUMERO DE DOCUMENTO', 'DIGITO DE VERIFICACION', 'TIPO DE PERDIDA', 'CLASE RIESGO', 'TIPO DE MOVIMIENTO', 'PRODUCTO', 'PROCESO', 'LINEA OPERATIVA', 'VALOR BASE']
class MEPParser(Parser):
 tipo_insumo="MEP"
 def parse(self,source:Path):
  r=IngestionResult(self.tipo_insumo); wb=load_workbook(source,data_only=True,read_only=True); ws=wb[wb.sheetnames[0]]
  hdr=[ws.cell(1,c).value for c in range(1,ws.max_column+1)]
  if hdr!=EXPECTED: r.errors.append(RecordError(1,None,"HEADER_INVALIDO","Encabezado MEP no coincide")); return r
  for n,vals in enumerate(ws.iter_rows(min_row=2,values_only=True),start=2):
   if any(v not in (None,"") for v in vals): r.records.append({"numero_linea":n,**dict(zip(EXPECTED,vals))})
  r.controls.append({"codigo":"TOTAL_REGISTROS","obtenido":len(r.records)}); return r
