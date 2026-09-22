from pathlib import Path
from openpyxl import load_workbook
from app.parsers.base import Parser
from app.ingestion.result import IngestionResult,RecordError
EXPECTED=['Tipo Transacción', 'Número de Tarjeta', 'Fecha de Comprobante', 'Código de Autorización', 'Número de Comprobante', 'Valor Pesos', 'Comisión Intercambio Pesos', 'Valor Dólares', 'Comisión Intercambio Dólares', 'Código Establecimiento', 'Código único', 'Nombre Establecimiento', 'Motivo de Contracargo', 'Referencia Contracargo', 'Referencia Universal', 'MCC', 'Ciudad', 'Indicador de Documentación', 'DE 72 – Texto', 'Datos Punto de Servicio', 'IRD']
class PMDParser(Parser):
 tipo_insumo="PMD"
 def parse(self,source:Path):
  r=IngestionResult(self.tipo_insumo); wb=load_workbook(source,data_only=True,read_only=True)
  if "PMD_Detalle" not in wb.sheetnames:
   r.errors.append(RecordError(None,None,"HOJA_NO_ENCONTRADA","PMD_Detalle")); return r
  ws=wb["PMD_Detalle"]; hdr=[ws.cell(1,c).value for c in range(1,ws.max_column+1)]
  if hdr!=EXPECTED:
   r.errors.append(RecordError(1,None,"HEADER_INVALIDO","Encabezado PMD_Detalle no coincide")); return r
  for n,vals in enumerate(ws.iter_rows(min_row=2,values_only=True),start=2):
   if any(v not in (None,"") for v in vals): r.records.append({"numero_linea":n,**dict(zip(EXPECTED,vals))})
  r.controls.append({"codigo":"TOTAL_REGISTROS","obtenido":len(r.records)}); return r
