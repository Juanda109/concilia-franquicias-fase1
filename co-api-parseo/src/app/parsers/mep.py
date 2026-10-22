from datetime import date
from pathlib import Path
from openpyxl import load_workbook
from app.parsers.base import Parser
from app.ingestion.result import IngestionResult,RecordError
EXPECTED=['CENTRO DE COSTO', 'CUENTA', 'DIVISA', 'CONTRATO', 'REFERENCIA DE CRUCE', 'IMPORTE', 'DESCRIPCION', 'FECHA', 'TIPO DE DOCUMENTO', 'NUMERO DE DOCUMENTO', 'DIGITO DE VERIFICACION', 'TIPO DE PERDIDA', 'CLASE RIESGO', 'TIPO DE MOVIMIENTO', 'PRODUCTO', 'PROCESO', 'LINEA OPERATIVA', 'VALOR BASE']
# EXPECTED es el encabezado literal del Excel; CON_MEP_RESULTADO
# (16_con_mep_resultado.sql) usa otros nombres de columna (sin "DE", con
# guion bajo) - el insert generico (resultados/generic.py) usa las claves
# del record tal cual como nombres de columna SQL, por eso hay que mapear.
COLUMNAS=['CENTRO_COSTO', 'CUENTA', 'DIVISA', 'CONTRATO', 'REFERENCIA_CRUCE', 'IMPORTE', 'DESCRIPCION', 'FECHA', 'TIPO_DOCUMENTO', 'NUMERO_DOCUMENTO', 'DIGITO_VERIFICACION', 'TIPO_PERDIDA', 'CLASE_RIESGO', 'TIPO_MOVIMIENTO', 'PRODUCTO', 'PROCESO', 'LINEA_OPERATIVA', 'VALOR_BASE']
def _fecha(v):
 """FECHA llega como texto AAAAMMDD (p.ej. "20260706"); FECHA es DATE en el
 DDL (16_con_mep_resultado.sql), hay que convertirla o Postgres rechaza el
 INSERT. Valores que no calcen van como NULL en vez de romper el archivo."""
 v=str(v) if v is not None else ""
 if len(v)!=8 or not v.isdigit(): return None
 try: return date(int(v[0:4]),int(v[4:6]),int(v[6:8])).isoformat()
 except ValueError: return None
class MEPParser(Parser):
 tipo_insumo="MEP"
 def parse(self,source:Path):
  r=IngestionResult(self.tipo_insumo); wb=load_workbook(source,data_only=True,read_only=True); ws=wb[wb.sheetnames[0]]
  hdr=[ws.cell(1,c).value for c in range(1,ws.max_column+1)]
  if hdr!=EXPECTED: r.errors.append(RecordError(1,None,"HEADER_INVALIDO","Encabezado MEP no coincide")); return r
  for n,vals in enumerate(ws.iter_rows(min_row=2,values_only=True),start=2):
   if any(v not in (None,"") for v in vals):
    row={"numero_linea":n,**dict(zip(COLUMNAS,vals))}
    row["FECHA"]=_fecha(row["FECHA"])
    r.records.append(row)
  r.controls.append({"codigo":"TOTAL_REGISTROS","obtenido":len(r.records)}); return r
