import re
from pathlib import Path
from openpyxl import load_workbook
from app.parsers.base import Parser
from app.ingestion.result import IngestionResult,RecordError

NUM=r'-?[\d.]+(?:,\d+)?-?'

def _num(tok):
 tok=tok.strip()
 if not tok: return None
 neg=tok.endswith('-')
 if neg: tok=tok[:-1]
 if tok.startswith('-'): neg=True; tok=tok[1:]
 tok=tok.replace('.','').replace(',','.')
 try: v=float(tok)
 except ValueError: return None
 return -v if neg else v

def _tokens(txt):
 return re.findall(NUM,txt)

RE_PRODUCTO=re.compile(r'^\s*([A-Za-zÁÉÍÓÚÑ0-9/.\- ]+?)\s{2,}((?:'+NUM+r'\s*){2,7})\s*$')
RE_TOTALES_FAVOR=re.compile(r'T\s*O\s*T\s*A\s*L\s*E\s*S\s+A\s+F\s*A\s*V\s*O\s*R',re.I)
RE_TOTALES_CARGO=re.compile(r'T\s*O\s*T\s*A\s*L\s*E\s*S\s+A\s+C\s*A\s*R\s*G\s*O',re.I)
RE_INTERNACIONAL_DETALLE=re.compile(r'^\s*(\d+)\s+(\d{4}-\d{2}-\d{2})\s+(\S+)\s+('+NUM+r')\s+('+NUM+r')\s+('+NUM+r')\s+([\d.,]+)\s+([\d.,]+)\s*$')
RE_INTERNACIONAL_TOTALES=re.compile(r'^\s*TOTALES\s+('+NUM+r')\s+('+NUM+r')\s+('+NUM+r')\s*$',re.I)
RE_POSICION_NETA_228=re.compile(r'POSICION\s+NETA\s*:\s*('+NUM+r')',re.I)
RE_CARGO_DE=re.compile(r'CARGO\s+DE\s+\$\s*('+NUM+r')',re.I)
RE_SUBTOTAL_NO_GRAVADO=re.compile(r'SUBTOTAL\s+NO\s+GRAVADO.*?:\s*('+NUM+r')',re.I)
RE_SUBTOTAL_GRAVADO=re.compile(r'SUBTOTAL\s+GRAVADO.*?:\s*('+NUM+r')',re.I)
RE_CONCEPTO_228=re.compile(r'^\s*(.+?)\s*:?\s*(?:A\s+(FAVOR|CARGO)\s+)?('+NUM+r')\s*$',re.I)
RE_SOLO_SEPARADOR=re.compile(r'^[\s\-]*$')

MARCADORES_PRESENTACION=[
 'FECHA DE IMPRESION','CAI0224A','CAI0228A','CARTA DE COMPENSACION','INFORME COMPLEMENTARIO',
 'ENTIDAD','FECHA:','SANTA FE DE BOGOTA','SE#ORES','SEÑORES','BANCO','CIUDAD','ESTIMADOS',
 'ADJUNTO LES ESTAMOS','CANJE ENVIADO, ASI','IGUALMENTE REMITO','PARA EFECTOS','CORDIALMENTE',
 'CREDIBANCO','C R E D I B A N C O','DIRECCION DE OPERACIONES','VISTO BUENO','REMITIMOS INFORMES','LA LIQUIDACION QUE ANEXAMOS',
 'VALORES CANJE INTERNACIONAL','CONCEPTOS RELACIONADOS','CONCEPTOS ASOCIADOS','VALORES NETOS',
 'PROCESO INTERBANCARIO','FECHA     FECHA','JULIANA','FACTOR/DOLAR','P R O D U C T O','V A L O R','REINTEGROS','DESCUENTOS',
 'PAGINA','FECHA :',
]

def _es_presentacion(norm):
 if RE_SOLO_SEPARADOR.match(norm): return True
 up=norm.upper()
 return any(m in up for m in MARCADORES_PRESENTACION)

class CartaCompensacionParser(Parser):
 tipo_insumo="CARTA_COMPENSACION"
 sheet_224="carta compensacion 224A"
 sheet_228="CAI228A"

 def parse(self,source:Path):
  r=IngestionResult(self.tipo_insumo)
  wb=load_workbook(source,data_only=True,read_only=True)
  if self.sheet_224 not in wb.sheetnames or self.sheet_228 not in wb.sheetnames:
   r.errors.append(RecordError(None,None,"HOJA_NO_ENCONTRADA",f"{self.sheet_224} / {self.sheet_228}")); return r

  descartes={}
  def _descarta(motivo):
   descartes[motivo]=descartes.get(motivo,0)+1

  posicion_neta_224=[]
  self._parse_224(wb[self.sheet_224],r,_descarta,posicion_neta_224)
  posicion_neta_228=[]
  self._parse_228(wb[self.sheet_228],r,_descarta,posicion_neta_228)

  if len(posicion_neta_224)==1 and len(posicion_neta_228)==1:
   v224,v228=posicion_neta_224[0],posicion_neta_228[0]
   if v224==v228:
    r.controls.append({"codigo":"CUADRE_POSICION_NETA","esperado":v228,"obtenido":v224,"estado":"OK"})
   else:
    r.errors.append(RecordError(None,None,"DESCUADRE_POSICION_NETA",
     f"Posición neta CAI0224A ({v224}) no coincide con CAI0228A ({v228})"))
  else:
   r.errors.append(RecordError(None,None,"DESCUADRE_POSICION_NETA",
    f"No se pudo determinar una única posición neta por hoja (224A={posicion_neta_224}, 228A={posicion_neta_228})"))

  for motivo,cantidad in descartes.items():
   r.controls.append({"codigo":f"DESCARTADOS_{motivo}","obtenido":cantidad})
  r.controls.append({"codigo":"TOTAL_REGISTROS","obtenido":len(r.records)})
  return r

 def _parse_224(self,ws,r,_descarta,posicion_neta_224):
  modo=None # None | 'ENVIADO' | 'RECIBIDO'
  esperando=None # None | 'REINTEGRO' | 'DESCUENTO'
  neta_vista=False
  for n in range(1,ws.max_row+1):
   line=str(ws.cell(n,1).value or "")
   norm=re.sub(r'\s+',' ',line).strip()
   if not norm: continue

   m=RE_INTERNACIONAL_DETALLE.match(line)
   if m:
    _,_fecha,cinta,outgoing,incoming,neto,fcompra,fventa=m.groups()
    outgoing_v,incoming_v,neto_v=_num(outgoing),_num(incoming),_num(neto)
    if outgoing_v in (0,None) and incoming_v in (0,None) and neto_v in (0,None):
     _descarta("CANJE_INTERNACIONAL_SIN_MOVIMIENTO"); continue
    r.records.append({"numero_linea":n,"seccion":"CAI0224A","motivo_codigo":"CANJE_INTERNACIONAL","es_control":False,
     "lado":None,"concepto":cinta,"valor_principal":neto_v,"valor_outgoing":outgoing_v,"valor_incoming":incoming_v,
     "factor_compra":_num(fcompra),"factor_venta":_num(fventa),"cinta":cinta,"fecha_cinta":_fecha})
    continue

   m=RE_INTERNACIONAL_TOTALES.match(line)
   if m:
    outgoing,incoming,neto=m.groups()
    r.records.append({"numero_linea":n,"seccion":"CAI0224A","motivo_codigo":"TOTAL_CANJE_INTERNACIONAL","es_control":True,
     "lado":None,"concepto":"TOTALES CANJE INTERNACIONAL","valor_principal":_num(neto),
     "valor_outgoing":_num(outgoing),"valor_incoming":_num(incoming)})
    continue

   if not neta_vista:
    m=RE_CARGO_DE.search(line)
    if m:
     valor=_num(m.group(1))
     r.records.append({"numero_linea":n,"seccion":"CAI0224A","motivo_codigo":"POSICION_NETA","es_control":True,
      "lado":None,"concepto":"VALOR NETO A SU CARGO","valor_principal":valor})
     posicion_neta_224.append(valor); neta_vista=True
     continue

   if RE_TOTALES_FAVOR.search(line):
    vals=_tokens(line)
    self._emitir_total_canje(r,n,"CAI0224A","A_FAVOR","TOTAL_CANJE_ENVIADO",vals)
    modo=None; continue
   if RE_TOTALES_CARGO.search(line):
    vals=_tokens(line)
    self._emitir_total_canje(r,n,"CAI0224A","A_CARGO","TOTAL_CANJE_RECIBIDO",vals)
    modo=None; continue

   up=norm.upper()
   if 'PRODUCTO' in up.replace(' ','') and ('AFAVOR' in up.replace(' ','')):
    modo='ENVIADO'; continue
   if 'PRODUCTO' in up.replace(' ','') and ('ACARGO' in up.replace(' ','')):
    modo='RECIBIDO'; continue

   if up=='REINTEGROS': esperando='REINTEGRO'; continue
   if up=='DESCUENTOS': esperando='DESCUENTO'; continue
   if esperando:
    vals=_tokens(line)
    if vals:
     valor=_num(vals[-1]); concepto=line[:line.rfind(vals[-1])].strip()
     if esperando=='REINTEGRO':
      if valor: r.records.append({"numero_linea":n,"seccion":"CAI0224A","motivo_codigo":"REINTEGRO","es_control":False,
       "lado":None,"concepto":concepto,"valor_principal":valor})
      else: _descarta("CONCEPTO_SIN_MOVIMIENTO")
     else:
      if valor: r.records.append({"numero_linea":n,"seccion":"CAI0224A","motivo_codigo":"DESCUENTO_CARGO","es_control":False,
       "lado":None,"concepto":concepto,"valor_principal":valor})
      else: _descarta("CONCEPTO_SIN_MOVIMIENTO")
     esperando=None; continue
    esperando=None

   if modo in ('ENVIADO','RECIBIDO'):
    m=RE_PRODUCTO.match(line)
    if m:
     nombre,valores_txt=m.groups(); vals=[_num(v) for v in _tokens(valores_txt)]
     if all((v is None or v==0) for v in vals):
      _descarta("CONCEPTO_SIN_MOVIMIENTO"); continue
     if modo=='ENVIADO' and len(vals)>=7:
      r.records.append({"numero_linea":n,"seccion":"CAI0224A","motivo_codigo":"CONCEPTO_CANJE_ENVIADO","es_control":False,
       "lado":"A_FAVOR","concepto":nombre.strip(),"valor_principal":vals[0],"valor_proceso":vals[1],
       "valor_financiacion":vals[2],"valor_fotocopias":vals[3],"valor_red":vals[4],
       "valor_depositos_electronicos":vals[5],"valor_total":vals[6]})
      continue
     if modo=='RECIBIDO' and len(vals)>=6:
      r.records.append({"numero_linea":n,"seccion":"CAI0224A","motivo_codigo":"CONCEPTO_CANJE_RECIBIDO","es_control":False,
       "lado":"A_CARGO","concepto":nombre.strip(),"valor_principal":vals[0],"valor_proceso":vals[1],
       "valor_financiacion":vals[2],"valor_fotocopias":vals[3],"valor_red":vals[4],"valor_total":vals[5]})
      continue

   if _es_presentacion(norm): continue
   _descarta("SIN_REGLA_FUNCIONAL_CARTA")

 def _emitir_total_canje(self,r,n,seccion,lado,motivo,vals):
  nums=[_num(v) for v in vals]
  fila={"numero_linea":n,"seccion":seccion,"motivo_codigo":motivo,"es_control":True,"lado":lado,
   "concepto":f"TOTALES {lado.replace('_',' ')}"}
  if lado=='A_FAVOR' and len(nums)>=7:
   fila.update(valor_principal=nums[0],valor_proceso=nums[1],valor_financiacion=nums[2],valor_fotocopias=nums[3],
    valor_red=nums[4],valor_depositos_electronicos=nums[5],valor_total=nums[6])
  elif lado=='A_CARGO' and len(nums)>=6:
   fila.update(valor_principal=nums[0],valor_proceso=nums[1],valor_financiacion=nums[2],valor_fotocopias=nums[3],
    valor_red=nums[4],valor_total=nums[5])
  r.records.append(fila)

 def _parse_228(self,ws,r,_descarta,posicion_neta_228):
  for n in range(1,ws.max_row+1):
   line=str(ws.cell(n,1).value or "")
   norm=re.sub(r'\s+',' ',line).strip()
   if not norm: continue
   if _es_presentacion(norm): continue

   m=RE_POSICION_NETA_228.search(line)
   if m:
    valor=_num(m.group(1))
    r.records.append({"numero_linea":n,"seccion":"CAI0228A","motivo_codigo":"POSICION_NETA","es_control":True,
     "lado":None,"concepto":"POSICION NETA","valor_principal":valor})
    posicion_neta_228.append(valor); continue

   m=RE_SUBTOTAL_NO_GRAVADO.search(line)
   if m:
    r.records.append({"numero_linea":n,"seccion":"CAI0228A","motivo_codigo":"SUBTOTAL_NO_GRAVADO","es_control":True,
     "lado":None,"concepto":"SUBTOTAL NO GRAVADO","valor_principal":_num(m.group(1))}); continue

   m=RE_SUBTOTAL_GRAVADO.search(line)
   if m:
    r.records.append({"numero_linea":n,"seccion":"CAI0228A","motivo_codigo":"SUBTOTAL_GRAVADO","es_control":True,
     "lado":None,"concepto":"SUBTOTAL GRAVADO","valor_principal":_num(m.group(1))}); continue

   m=RE_CONCEPTO_228.match(line)
   if m:
    concepto,lado,valor_txt=m.groups(); valor=_num(valor_txt)
    lado_norm=('A_FAVOR' if lado.upper()=='FAVOR' else 'A_CARGO') if lado else None
    if not valor:
     _descarta("CONTROL_SIN_MOVIMIENTO"); continue
    r.records.append({"numero_linea":n,"seccion":"CAI0228A","motivo_codigo":"CONCEPTO_CONTROL_CAI228A","es_control":True,
     "lado":lado_norm,"concepto":concepto.strip(),"valor_principal":valor})
    continue

   _descarta("SIN_REGLA_FUNCIONAL_CARTA")
