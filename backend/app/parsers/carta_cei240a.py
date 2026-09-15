from pathlib import Path
from openpyxl import load_workbook
from app.parsers.base import Parser
from app.ingestion.result import IngestionResult,RecordError
import re
class CartaCEI240AParser(Parser):
 tipo_insumo="CARTA_COMPENSACION"; sheet="CEI  240A"
 def parse(self,source:Path):
  r=IngestionResult(self.tipo_insumo); wb=load_workbook(source,data_only=True,read_only=True)
  if self.sheet not in wb.sheetnames: r.errors.append(RecordError(None,None,"HOJA_NO_ENCONTRADA",self.sheet)); return r
  ws=wb[self.sheet]; lines=[str(ws.cell(i,1).value or "").rstrip() for i in range(1,ws.max_row+1)]
  if "CEI0240A" not in "\n".join(lines): r.errors.append(RecordError(None,None,"REPORTE_INVALIDO","CEI0240A no encontrado")); return r
  pattern=re.compile(r"^\s*0\s+(?:(\d{2})\s+)?(VENTAS|REVERSION VENTAS|DISPENSACION EFECTIVO EN A\.T\.M\.|CONSULTAS SALDO EN A\.T\.M|TRANSACCIONES DECLINADAS EN A\.T\.M\.|TRANSFERENCIA ENTRE CUENTAS POR A\.T\.M\.)",re.I)
  for n,line in enumerate(lines,start=1):
   m=pattern.search(line)
   if m:r.records.append({"numero_linea":n,"codigo_movimiento":m.group(1),"concepto":m.group(2).upper()})
  for n,line in enumerate(lines,start=1):
   if "TOTALES GENERALES" in line.upper():r.records.append({"numero_linea":n,"codigo_movimiento":None,"concepto":"TOTALES GENERALES"})
  r.controls.append({"codigo":"BLOQUES_CEI240A","obtenido":len(r.records)}); return r
