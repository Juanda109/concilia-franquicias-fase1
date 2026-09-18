"""Genera el cuaderno de evidencias funcionales en .docx, con el formato del banco.

Lee ``docs/EVIDENCIAS_TNR_RCS.md``, hereda los estilos de la plantilla que ya usa
el equipo (``Evidencias_TXNR_DEV_v1.docx``) e inserta las capturas tratadas de
``docs/evidencias_rcs/`` donde el markdown las referencia.

    uv run --with python-docx python scripts/build_evidencias_docx.py
    uv run --with python-docx python scripts/build_evidencias_docx.py --md docs/OTRO.md

El markdown sigue siendo la fuente. Este script solo lo viste con el formato de
entrega: si hay que cambiar el contenido, se cambia el .md y se regenera.
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
MD = ROOT / "docs" / "EVIDENCIAS_TNR_RCS.md"
IMAGENES = ROOT / "docs" / "evidencias_rcs"
PLANTILLA = (
    ROOT
    / "co_pqrs_back_trx_noreconocida"
    / "docs"
    / "report"
    / "Evidencias_TXNR_DEV_v1.docx"
)
SALIDA = ROOT / "docs" / "EVIDENCIAS_TNR_RCS.docx"

ANCHO_IMAGEN = Inches(6.3)
ESTILO_TABLA = "Light Grid Accent 1"

RE_ENCABEZADO = re.compile(r"^(#{1,6})\s+(.*)$")
RE_IMAGEN = re.compile(r"^`([^`]+\.png)`$")
RE_PIE = re.compile(r"^\*\*(T\d{2}[a-z]):\*\*\s*(.*)$")
# Trozos con formato: **negrita** o `codigo`. El resto va tal cual.
RE_TROZOS = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")


def documento_en_blanco() -> Document:
    """Abre la plantilla del banco y vacia su contenido, conservando estilos.

    Se conserva ``sectPr``, que lleva el tamano de pagina y los margenes.
    """

    doc = Document(str(PLANTILLA))
    cuerpo = doc.element.body
    for hijo in list(cuerpo):
        if hijo.tag.endswith("}sectPr"):
            continue
        cuerpo.remove(hijo)
    return doc


def escribir_texto(parrafo, texto: str) -> None:
    """Vuelca ``texto`` en el parrafo resolviendo negritas y codigo."""

    for trozo in RE_TROZOS.split(texto):
        if not trozo:
            continue
        if trozo.startswith("**") and trozo.endswith("**"):
            parrafo.add_run(trozo[2:-2]).bold = True
        elif trozo.startswith("`") and trozo.endswith("`"):
            run = parrafo.add_run(trozo[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        else:
            parrafo.add_run(trozo)


def celdas(linea: str) -> list[str]:
    return [c.strip() for c in linea.strip().strip("|").split("|")]


def es_separador(linea: str) -> bool:
    return bool(re.fullmatch(r"\|[\s:|-]+\|", linea.strip()))


def anadir_tabla(doc: Document, filas: list[str]) -> None:
    """Escribe una tabla markdown. La primera fila es la cabecera."""

    cabecera = celdas(filas[0])
    cuerpo = [celdas(f) for f in filas[2:]]
    # Las fichas de datos (la de portada, por ejemplo) van sin cabecera: en el
    # markdown se escriben con la primera fila vacia, y una franja sombreada sin
    # texto quedaria fea en Word.
    sin_cabecera = not any(cabecera)
    tabla = doc.add_table(rows=0 if sin_cabecera else 1, cols=len(cabecera))
    try:
        tabla.style = ESTILO_TABLA
    except KeyError:  # la plantilla siempre lo trae, pero no cuesta ser prudente
        pass
    if not sin_cabecera:
        for celda, texto in zip(tabla.rows[0].cells, cabecera):
            celda.paragraphs[0].text = ""
            escribir_texto(celda.paragraphs[0], texto)
    for fila in cuerpo:
        nuevas = tabla.add_row().cells
        for celda, texto in zip(nuevas, fila):
            celda.paragraphs[0].text = ""
            escribir_texto(celda.paragraphs[0], texto)
    if sin_cabecera and len(cabecera) == 2:
        # Ficha de dos columnas: etiqueta estrecha, contenido ancho. Word respeta
        # el ancho de celda cuando se fija en todas las filas.
        tabla.autofit = False
        for fila_t in tabla.rows:
            fila_t.cells[0].width = Inches(1.55)
            fila_t.cells[1].width = Inches(4.75)
    doc.add_paragraph()


def anadir_cita(doc: Document, lineas: list[str]) -> None:
    parrafo = doc.add_paragraph()
    parrafo.paragraph_format.left_indent = Inches(0.3)
    escribir_texto(parrafo, " ".join(lineas))
    for run in parrafo.runs:
        run.italic = True


def anadir_imagen(doc: Document, nombre: str) -> bool:
    ruta = IMAGENES / nombre
    if not ruta.exists():
        print(f"  AVISO: falta la captura {nombre}, se deja un hueco senalado")
        hueco = doc.add_paragraph()
        hueco.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = hueco.add_run(f"[pendiente de capturar: {nombre}]")
        run.italic = True
        return False
    doc.add_picture(str(ruta), width=ANCHO_IMAGEN)
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    return True


def anadir_codigo(doc: Document, lineas_codigo: list[str]) -> None:
    """Vuelca un aparte de registro conservando sus saltos de linea."""

    parrafo = doc.add_paragraph()
    parrafo.paragraph_format.left_indent = Inches(0.25)
    parrafo.paragraph_format.space_after = Pt(8)
    for n, linea in enumerate(lineas_codigo):
        run = parrafo.add_run(linea)
        run.font.name = "Consolas"
        run.font.size = Pt(8)
        if n < len(lineas_codigo) - 1:
            run.add_break()


def anadir_pie(doc: Document, codigo: str, texto: str) -> None:
    parrafo = doc.add_paragraph()
    parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    etiqueta = parrafo.add_run(f"{codigo}: ")
    etiqueta.bold = True
    etiqueta.font.size = Pt(9)
    resto = parrafo.add_run(texto)
    resto.italic = True
    resto.font.size = Pt(9)


def main() -> None:
    global MD, SALIDA
    import argparse

    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--md", type=Path, default=MD, help="fuente markdown (por defecto el cuaderno de evidencias)")
    ap.add_argument("--out", type=Path, default=None, help="docx de salida (por defecto, junto a la fuente)")
    args = ap.parse_args()
    MD = args.md.resolve()
    SALIDA = (args.out or MD.with_suffix(".docx")).resolve()

    if not MD.exists():
        sys.exit(f"No encuentro {MD}")
    if not PLANTILLA.exists():
        sys.exit(f"No encuentro la plantilla {PLANTILLA}")

    doc = documento_en_blanco()
    lineas = MD.read_text(encoding="utf-8").split("\n")

    i = 0
    n_imagenes = 0
    primer_encabezado = True
    parrafo_pendiente: list[str] = []

    def cerrar_parrafo() -> None:
        if parrafo_pendiente:
            escribir_texto(doc.add_paragraph(), " ".join(parrafo_pendiente))
            parrafo_pendiente.clear()

    while i < len(lineas):
        linea = lineas[i]
        desnuda = linea.strip()

        if not desnuda:
            cerrar_parrafo()
            i += 1
            continue

        if desnuda == "---":
            cerrar_parrafo()
            i += 1
            continue

        m = RE_ENCABEZADO.match(desnuda)
        if m:
            cerrar_parrafo()
            nivel, texto = len(m.group(1)), m.group(2)
            if primer_encabezado and nivel == 1:
                doc.add_heading(texto, level=0)
                primer_encabezado = False
            elif nivel == 2 and len(doc.paragraphs) <= 2:  # subtitulo de portada
                parrafo = doc.add_paragraph()
                run = parrafo.add_run(texto)
                run.bold = True
                run.font.size = Pt(13)
            else:
                if nivel == 2 and any(p.style.name == "Heading 1" for p in doc.paragraphs):
                    doc.add_page_break()
                doc.add_heading(texto, level=min(nivel - 1, 4))
            i += 1
            continue

        if desnuda.startswith("```"):
            cerrar_parrafo()
            i += 1
            bloque = []
            while i < len(lineas) and not lineas[i].strip().startswith("```"):
                bloque.append(lineas[i].rstrip())
                i += 1
            i += 1  # la linea de cierre
            anadir_codigo(doc, bloque)
            continue

        m = RE_IMAGEN.match(desnuda)
        if m:
            cerrar_parrafo()
            if anadir_imagen(doc, m.group(1)):
                n_imagenes += 1
            i += 1
            continue

        m = RE_PIE.match(desnuda)
        if m:
            cerrar_parrafo()
            anadir_pie(doc, m.group(1), m.group(2))
            i += 1
            continue

        if desnuda.startswith("|"):
            cerrar_parrafo()
            bloque = []
            while i < len(lineas) and lineas[i].strip().startswith("|"):
                bloque.append(lineas[i].strip())
                i += 1
            if len(bloque) >= 2 and es_separador(bloque[1]):
                anadir_tabla(doc, bloque)
            else:  # tabla mal formada: se vuelca tal cual para no perder contenido
                for fila in bloque:
                    escribir_texto(doc.add_paragraph(), fila)
            continue

        if desnuda.startswith(">"):
            cerrar_parrafo()
            bloque = []
            while i < len(lineas) and lineas[i].strip().startswith(">"):
                bloque.append(lineas[i].strip().lstrip(">").strip())
                i += 1
            anadir_cita(doc, bloque)
            continue

        if re.match(r"^\d+\.\s+|^[-*]\s+", desnuda):
            cerrar_parrafo()
            texto = re.sub(r"^\d+\.\s+|^[-*]\s+", "", desnuda)
            estilo = "List Number" if re.match(r"^\d+\.", desnuda) else "List Bullet"
            try:
                parrafo = doc.add_paragraph(style=estilo)
            except KeyError:
                parrafo = doc.add_paragraph()
            escribir_texto(parrafo, texto)
            i += 1
            continue

        parrafo_pendiente.append(desnuda)
        i += 1

    cerrar_parrafo()
    doc.save(str(SALIDA))
    try:
        mostrado = SALIDA.relative_to(ROOT)
    except ValueError:  # salida fuera del repositorio: se muestra completa
        mostrado = SALIDA
    print(f"  escrito {mostrado}")
    print(f"  {len(doc.paragraphs)} parrafos, {len(doc.tables)} tablas, {n_imagenes} capturas")


if __name__ == "__main__":
    main()
