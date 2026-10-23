# Autor: Gustavo Andres Camargo Duque
# Identificador: C805925
# Proyecto: Conciliacion Franquicias

"""Parser F1 PMDGeneralEntidadAAMMDD / AAAAMMDD (.xls/.xlsx).
Mantiene el contrato app.parsers.base / app.ingestion.result.
Los niveles son identificadores jerarquicos, no numeros decimales.
No asigna cuentas contables ni deduce reglas F2 sin homologacion.
"""
from __future__ import annotations
from decimal import Decimal, InvalidOperation
from pathlib import Path
import re
from openpyxl import load_workbook
from app.parsers.base import Parser
from app.ingestion.result import IngestionResult, RecordError

COLUMNS = (
    'PRODUCTO_SERVICIO','NIVEL',
    'DEBITO_CANTIDAD','DEBITO_VALOR','DEBITO_COMISION_INTERCAMBIO',
    'DEBITO_COMISION_ADMINISTRATIVA','DEBITO_NETO',
    'CREDITO_CANTIDAD','CREDITO_VALOR','CREDITO_COMISION_INTERCAMBIO',
    'CREDITO_COMISION_ADMINISTRATIVA','CREDITO_NETO','TOTAL',
)
NUMERIC = COLUMNS[2:]
NAME_RE = re.compile(r'^PMDGeneralEntidad(\d{6}|\d{8})(?:\s*\(\d+\))?\.(xls|xlsx)$', re.I)
LEVEL_RE = re.compile(r'^\d+(?:\.\d+)*$')


def _decimal(value):
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, bool):
        raise ValueError('booleano no valido')
    try:
        d = Decimal(str(value).strip().replace(',', ''))
    except InvalidOperation as exc:
        raise ValueError(f'valor numerico invalido: {value!r}') from exc
    if not d.is_finite():
        raise ValueError('valor no finito')
    return d


def _open_xls(path):
    try:
        import xlrd
    except ImportError as exc:
        raise RuntimeError('Los archivos .xls requieren xlrd>=2.0.1; instale la dependencia.') from exc
    book = xlrd.open_workbook(str(path), on_demand=True)
    if 'PMD' not in book.sheet_names():
        book.release_resources()
        raise ValueError('HOJA_NO_ENCONTRADA: PMD')
    sheet = book.sheet_by_name('PMD')
    try:
        for i in range(sheet.nrows):
            yield i + 1, [sheet.cell_value(i, j) if j < sheet.ncols else None for j in range(13)]
    finally:
        book.release_resources()


def _open_xlsx(path):
    book = load_workbook(path, data_only=True, read_only=True)
    try:
        if 'PMD' not in book.sheetnames:
            raise ValueError('HOJA_NO_ENCONTRADA: PMD')
        for i, row in enumerate(book['PMD'].iter_rows(values_only=True), 1):
            yield i, list(row[:13]) + [None] * max(0, 13 - len(row))
    finally:
        book.close()


class PMDParser(Parser):
    tipo_insumo = 'PMD'
    expected_columns = 13

    def parse(self, source: Path):
        source = Path(source)
        result = IngestionResult(self.tipo_insumo)
        m = NAME_RE.fullmatch(source.name)
        if not m:
            result.errors.append(RecordError(None, None, 'NOMBRE_INVALIDO',
                'Se espera PMDGeneralEntidadAAMMDD o PMDGeneralEntidadAAAAMMDD (.xls/.xlsx)'))
            return result
        if source.suffix.lower() not in ('.xls', '.xlsx'):
            result.errors.append(RecordError(None,None,'FORMATO_INVALIDO',source.suffix))
            return result
        try:
            rows = _open_xls(source) if source.suffix.lower()=='.xls' else _open_xlsx(source)
            count = 0
            for n, values in rows:
                count += 1
                if n == 1:
                    if str(values[0] or '').strip().upper() != 'PRODUCTO / SERVICIO' or str(values[1] or '').strip().upper() != 'NIVEL':
                        result.errors.append(RecordError(n,None,'HEADER_INVALIDO','Encabezado PMD inesperado'))
                        return result
                    continue
                if n == 2:
                    if str(values[2] or '').strip().lower() != 'cantidad' or str(values[7] or '').strip().lower() != 'cantidad':
                        result.errors.append(RecordError(n,None,'HEADER_INVALIDO','Segundo encabezado PMD inesperado'))
                        return result
                    continue
                if not any(v is not None and str(v).strip() for v in values):
                    continue
                producto = str(values[0] or '').strip()
                nivel = str(values[1] or '').strip()
                if not producto or not LEVEL_RE.fullmatch(nivel):
                    result.errors.append(RecordError(n,'NIVEL','REGISTRO_INVALIDO',
                        f'Producto o nivel invalido: {producto!r}, {nivel!r}'))
                    continue
                record = {'numero_linea':n, 'PRODUCTO_SERVICIO':producto, 'NIVEL':nivel,
                          'ES_CONTROL':producto.upper().startswith('TOTAL PRODUCTOS PMD'),
                          'FECHA_ARCHIVO':m.group(1)}
                try:
                    for key, value in zip(NUMERIC,values[2:]):
                        record[key] = _decimal(value)
                    for key in ('DEBITO_CANTIDAD','CREDITO_CANTIDAD'):
                        val = record[key]
                        if val is not None and (val < 0 or val != val.to_integral_value()):
                            raise ValueError(f'{key}: cantidad invalida {val}')
                except ValueError as exc:
                    result.errors.append(RecordError(n,None,'NUMERO_INVALIDO',str(exc)))
                    continue
                result.records.append(record)
            if count < 3:
                result.errors.append(RecordError(None,None,'ARCHIVO_VACIO','PMD sin filas de datos'))
        except (ValueError,RuntimeError,OSError,KeyError) as exc:
            result.errors.append(RecordError(None,None,'ARCHIVO_INVALIDO',str(exc)))
        result.controls.append({'codigo':'TOTAL_REGISTROS','obtenido':len(result.records)})
        result.controls.append({'codigo':'REGISTROS_CONTROL','obtenido':sum(r['ES_CONTROL'] for r in result.records)})
        return result
