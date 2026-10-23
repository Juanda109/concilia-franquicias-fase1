# Autor: Gustavo Andres Camargo Duque
# Identificador: C805925
# Proyecto: Conciliacion Franquicias

"""Parser F1 MEP: layout de 18 columnas observado en ANEXO 0017 MEP_COMPENSACION.
Retiene importes como Decimal y codigos/cuentas como texto.
No inventa clasificacion contable ni asiento automatico.
"""
from __future__ import annotations
from datetime import datetime, date
from decimal import Decimal, InvalidOperation
from pathlib import Path
from openpyxl import load_workbook
from app.parsers.base import Parser
from app.ingestion.result import IngestionResult, RecordError

COLUMNS = (
 'CENTRO DE COSTO','CUENTA','DIVISA','CONTRATO','REFERENCIA DE CRUCE',
 'IMPORTE','DESCRIPCION','FECHA','TIPO DE DOCUMENTO','NUMERO DE DOCUMENTO',
 'DIGITO DE VERIFICACION','TIPO DE PERDIDA','CLASE RIESGO',
 'TIPO DE MOVIMIENTO','PRODUCTO','PROCESO','LINEA OPERATIVA','VALOR BASE'
)

def _text(value):
    if value is None: return None
    if isinstance(value,float) and value.is_integer(): return str(int(value))
    return str(value).strip() or None

def _amount(value):
    if value is None or str(value).strip()=='':return None
    try: dec=Decimal(str(value).strip().replace(',',''))
    except InvalidOperation as exc: raise ValueError(f'importe invalido: {value!r}') from exc
    if not dec.is_finite(): raise ValueError('importe no finito')
    return dec

def _date(value):
    if value is None or str(value).strip()=='': return None
    if isinstance(value,(date,datetime)): return value.date().isoformat() if isinstance(value,datetime) else value.isoformat()
    raw=_text(value)
    for fmt in ('%Y%m%d','%Y-%m-%d'):
        try:return datetime.strptime(raw,fmt).date().isoformat()
        except ValueError: pass
    raise ValueError(f'fecha MEP invalida: {value!r}')

class MEPParser(Parser):
    tipo_insumo='MEP'
    expected_columns=18
    def parse(self,source:Path):
        result=IngestionResult(self.tipo_insumo)
        source=Path(source)
        if source.suffix.lower()!='.xlsx':
            result.errors.append(RecordError(None,None,'FORMATO_INVALIDO','Se espera libro .xlsx'))
            return result
        try:
            wb=load_workbook(source,data_only=True,read_only=True)
            try:
                if len(wb.sheetnames)!=1:
                    result.errors.append(RecordError(None,None,'HOJAS_INESPERADAS',str(wb.sheetnames)))
                    return result
                sheet=wb.active
                if sheet.max_column!=18:
                    result.errors.append(RecordError(1,None,'COLUMNAS_INVALIDAS',f'Se esperaban 18; hay {sheet.max_column}'))
                    return result
                iterator=sheet.iter_rows(values_only=True)
                header=next(iterator,None)
                if header is None or tuple(_text(v) or '' for v in header[:18])!=COLUMNS:
                    result.errors.append(RecordError(1,None,'HEADER_INVALIDO','Encabezado MEP no coincide con las 18 columnas esperadas'))
                    return result
                for n,values in enumerate(iterator,2):
                    if not any(v is not None and str(v).strip() for v in values):continue
                    record={'numero_linea':n}
                    record.update({key:_text(value) for key,value in zip(COLUMNS,values)})
                    try:
                        record['IMPORTE']=_amount(values[5]);record['VALOR BASE']=_amount(values[17])
                        record['FECHA']=_date(values[7])
                        if not record['CUENTA'] or record['IMPORTE'] is None or not record['FECHA']:
                            raise ValueError('CUENTA, IMPORTE y FECHA son obligatorios')
                    except ValueError as exc:
                        result.errors.append(RecordError(n,None,'REGISTRO_INVALIDO',str(exc)))
                        continue
                    result.records.append(record)
            finally:wb.close()
        except (OSError,ValueError,KeyError,StopIteration) as exc:
            result.errors.append(RecordError(None,None,'ARCHIVO_INVALIDO',str(exc)))
        result.controls.append({'codigo':'TOTAL_REGISTROS','obtenido':len(result.records)})
        result.controls.append({'codigo':'SUMA_IMPORTE','obtenido':sum((r['IMPORTE'] for r in result.records),Decimal('0'))})
        return result
